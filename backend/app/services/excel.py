"""Excel history of deals: restaurant, item, final price, distance, feedback, platform, last seen.

Every hunt (web or CLI) merges its deals into this file, keeping the Feedback you typed in Excel.
It is a history log; the web UI reads from the database.
"""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import SHEET_FILE
from app.services.pricing import Deal

HEADERS = ["Restaurant", "Item", "Final Price (₹)", "Distance (km)", "Feedback", "Platform", "Last seen"]
WIDTHS = [32, 42, 16, 14, 40, 12, 20]
COLS = len(HEADERS)


def _key(restaurant, item, price) -> tuple:
    return (str(restaurant or "").strip().lower(), str(item or "").strip().lower(), round(float(price or 0)))


def load_rows(path: Path = SHEET_FILE) -> list[list]:
    """Return existing data rows (without header); [] if the file is missing or unreadable."""
    if not path.exists():
        return []
    try:
        ws = load_workbook(path).active
        rows = [list(r[:COLS]) + [None] * (COLS - len(r)) for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]]
    except Exception as e:  # corrupt or open in Excel
        print(f"Could not read {path.name}: {e!r} -> starting fresh")
        return []
    return rows


def save_deals(deals: list[Deal], path: Path = SHEET_FILE) -> int:
    """Merge new deals into the existing sheet (keeping feedback), sort by price, rewrite the file."""
    rows = {_key(r[0], r[1], r[2]): r for r in load_rows(path)}
    for d in deals:
        k = _key(d.restaurant, d.items, d.final_payable)
        old = rows.get(k)
        rows[k] = [d.restaurant, d.items, d.final_payable, d.distance_km, old[4] if old else None, d.platform, d.found_at]
    data = sorted(rows.values(), key=lambda r: float(r[2] or 0))

    wb = Workbook()
    ws = wb.active
    ws.title = "Deals"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="D9480F")
    for r in data:
        ws.append(r)
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    try:
        wb.save(path)
    except PermissionError:
        print(f"Close {path.name} in Excel and run again - it is locked.")
        return 0
    return len(data)
